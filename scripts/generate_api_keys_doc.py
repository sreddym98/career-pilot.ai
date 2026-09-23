"""One-off generator for docs/API_KEYS_SETUP.docx and .pdf — run manually when
the key list changes: python scripts/generate_api_keys_doc.py"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
os.makedirs(OUT_DIR, exist_ok=True)

SECTIONS = [
    ("Free", [
        ("USAJOBS_EMAIL + USAJOBS_KEY", "https://developer.usajobs.gov/APIRequest/Index",
         "~40-80 federal QA/SDET roles"),
        ("ADZUNA_APP_ID + ADZUNA_APP_KEY", "https://developer.adzuna.com/",
         "~40-80 mixed board roles (1,000 calls/mo free)"),
        ("SUPABASE_URL + SUPABASE_JWT_SECRET", "https://supabase.com -> New project -> Settings -> API",
         "Real per-user accounts (without it everyone shares one dev user)"),
        ("GROQ_API_KEY", "https://console.groq.com/keys",
         "Low-cost AI generation for Autopilot (already set)"),
        ("GMAIL_CLIENT_ID + GMAIL_CLIENT_SECRET", "https://console.cloud.google.com/apis/credentials",
         "Gmail-based autopilot email sending/reading"),
        ("INTEGRATION_ENCRYPTION_KEY", "generate locally with Python's cryptography.fernet",
         "Encrypts stored integration tokens (already set)"),
        ("AUTH_SECRET", "generate locally with Python's secrets module",
         "Signs your own session tokens (required before production)"),
        ("STRIPE_PUBLISHABLE_KEY", "https://dashboard.stripe.com/apikeys",
         "Safe to expose client-side; pairs with the secret key below"),
    ]),
    ("Paid", [
        ("RAPIDAPI_KEY (JSearch)", "https://rapidapi.com/letscrape-6bRBa3QguO5/api/jsearch",
         "~$30/mo - Dice, LinkedIn, staffing/contract roles (already set, quota exhausted, upgrade plan)"),
        ("CORESIGNAL_API_KEY", "https://coresignal.com/pricing",
         "Usage-based - independent 482M+ global job postings dataset, second source beyond JSearch/Dice/LinkedIn"),
        ("ANTHROPIC_API_KEY", "https://console.anthropic.com -> API keys",
         "~$0.02-0.05/generation - resume tailoring + cover letters"),
        ("TWILIO_ACCOUNT_SID + TWILIO_AUTH_TOKEN + TWILIO_VERIFY_SERVICE_SID", "https://console.twilio.com",
         "Pay-as-you-go SMS - phone/SMS verification for Autopilot"),
    ]),
    ("Payments (Stripe)", [
        ("STRIPE_SECRET_KEY", "https://dashboard.stripe.com/apikeys", "Core Stripe secret key"),
        ("STRIPE_WEBHOOK_SECRET", "stripe listen --forward-to localhost:8000/api/billing/webhook",
         "Verifies webhook signatures"),
        ("STRIPE_PRICE_PRO_MONTHLY / RECRUITER / PRO_3MO / PRO_6MO / EVAL",
         "Printed automatically by running python setup_stripe.py",
         "Product price IDs"),
    ]),
    ("Infrastructure", [
        ("DATABASE_URL", "n/a - config only", "Defaults to local SQLite; switch to Postgres for real users"),
        ("REDIS_URL", "n/a - config only", "Only matters at scale (thousands of concurrent users)"),
        ("FRONTEND_URL", "n/a - config only", "Your deployed domain - required for CORS + Stripe redirects"),
        ("GMAIL_REDIRECT_URI", "n/a - config only", "Must exactly match what's registered in Google Cloud Console"),
        ("ENV", "n/a - config only", '"dev" locally, "prod" when deployed'),
    ]),
]

TITLE = "Career Pilot AI — API Keys & Setup Reference"
INTRO = ("Every environment variable this codebase reads, grouped by cost, "
         "with where to get it and what it unlocks. Verified against "
         "server/api/settings.py and server/ingest/sources.py.")

CURRENT_STATE = ("Already filled in server/.env: RAPIDAPI_KEY, GROQ_API_KEY, "
                  "STRIPE_SECRET_KEY, STRIPE_PUBLISHABLE_KEY, INTEGRATION_ENCRYPTION_KEY. "
                  "Still empty: USAJOBS_*, ADZUNA_*, ANTHROPIC_API_KEY, SUPABASE_*, "
                  "AUTH_SECRET, TWILIO_*, STRIPE_WEBHOOK_SECRET, and the Stripe price IDs.")


def build_docx(path):
    doc = Document()
    h = doc.add_heading(TITLE, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = doc.add_paragraph(INTRO)
    p.runs[0].italic = True

    note = doc.add_paragraph()
    note_run = note.add_run(CURRENT_STATE)
    note_run.bold = True
    note_run.font.color.rgb = RGBColor(0x0F, 0x66, 0x2A)

    for section_title, rows in SECTIONS:
        doc.add_heading(section_title, level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Key", "Where to get it", "Unlocks"
        for cell in hdr:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
        for key, where, unlock in rows:
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = key, where, unlock
        doc.add_paragraph()

    for section in doc.sections:
        section.top_margin = Pt(50)
        section.bottom_margin = Pt(50)
    doc.save(path)


def build_pdf(path):
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=18, spaceAfter=10)
    h2_style = ParagraphStyle("H2X", parent=styles["Heading2"], spaceBefore=16, spaceAfter=6,
                               textColor=colors.HexColor("#0958c7"))
    body_style = ParagraphStyle("BodyX", parent=styles["BodyText"], fontSize=9.5, leading=13)
    note_style = ParagraphStyle("NoteX", parent=styles["BodyText"], fontSize=9.5, leading=13,
                                 textColor=colors.HexColor("#0f662a"))

    doc = SimpleDocTemplate(path, pagesize=LETTER,
                             leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                             topMargin=0.6 * inch, bottomMargin=0.6 * inch)
    story = [Paragraph(TITLE, title_style),
             Paragraph(INTRO, body_style),
             Spacer(1, 8),
             Paragraph(CURRENT_STATE, note_style),
             Spacer(1, 10)]

    for section_title, rows in SECTIONS:
        story.append(Paragraph(section_title, h2_style))
        data = [["Key", "Where to get it", "Unlocks"]]
        for key, where, unlock in rows:
            data.append([Paragraph(key, body_style), Paragraph(where, body_style), Paragraph(unlock, body_style)])
        table = Table(data, colWidths=[2.0 * inch, 2.6 * inch, 2.4 * inch], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0958c7")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(table)
        story.append(Spacer(1, 10))

    doc.build(story)


if __name__ == "__main__":
    docx_path = os.path.join(OUT_DIR, "API_KEYS_SETUP.docx")
    pdf_path = os.path.join(OUT_DIR, "API_KEYS_SETUP.pdf")
    build_docx(docx_path)
    build_pdf(pdf_path)
    print(f"Wrote {os.path.abspath(docx_path)}")
    print(f"Wrote {os.path.abspath(pdf_path)}")
